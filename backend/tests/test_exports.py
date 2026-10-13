import io
import math
from decimal import Decimal

import httpx
import pypdfium2 as pdfium
import pytest
from openpyxl import load_workbook

from app.engine.units import LengthUnit, to_mm
from app.exports.data import NOTICE, NotReady, export_data
from app.exports.excel import FIRST_ROW, render_workbook
from app.exports.formatting import excel_money_format, format_length, format_money
from app.exports.pdf import render_report, render_shopping_list
from app.exports.shopping import shopping_html, shopping_text
from app.schemas.cost import CostLine, CostSettings, LineBasis
from app.schemas.material import MaterialOverride, Price, PriceBasis, ProjectMaterial, TileRef
from app.schemas.project import Room
from app.services import email

FT = LengthUnit.FT
T600 = TileRef(
    id="t6", name="Ivory matt", size_label="600 × 600 mm", length_mm=600, width_mm=600,
    tiles_per_box=4,
)  # fmt: skip
T300 = TileRef(
    id="t3", name="Grey anti-skid", size_label="300 × 300 mm", length_mm=300, width_mm=300,
    tiles_per_box=11,
)  # fmt: skip


def ft(n) -> Decimal:
    return to_mm(str(n), FT)


def room(name, **kw) -> Room:
    return Room(**({"id": name, "name": name, "confirmed": True, "unit": FT} | kw))


def project(settings: CostSettings | None = None, price: Price | None = None, rooms=None) -> dict:
    rooms = rooms or [
        room("Bedroom", length_mm=ft(11), width_mm=ft("9.5"), floor="First"),
        room("Kitchen", length_mm=ft(10), width_mm=ft(8)),
        room(
            "Bath & toilet",
            length_mm=ft(6),
            width_mm=ft(5),
            material=MaterialOverride(tile=T300, price=Price(amount=Decimal(60), basis="sqft")),
        ),
    ]
    default = ProjectMaterial(tile=T600, breakage_pct=Decimal(5), price=price)
    return {
        "name": "Sharma résidence",
        "rooms": [r.model_dump() for r in rooms],
        "default_material": default.model_dump(),
        "cost_settings": (settings or CostSettings()).model_dump(),
    }


SETTINGS = CostSettings(
    tax_label="GST",
    tax_pct=Decimal(18),
    budget=Decimal(1000),  # far below the total: over budget
    lines=[
        CostLine(label="Delivery", amount=Decimal(1500)),
        CostLine(label="Fixing", amount=Decimal(40), basis=LineBasis.SQFT, taxable=False),
        CostLine(label="Handling", amount=Decimal(10), basis=LineBasis.BOX),
    ],
)
PRICED = project(SETTINGS, Price(amount=Decimal(800), basis=PriceBasis.BOX))
UNPRICED = project(
    rooms=[
        room("Bedroom", length_mm=ft(11), width_mm=ft("9.5")),
        room("Bath", length_mm=ft(6), width_mm=ft(5), material=MaterialOverride(tile=T300)),
    ]
)


def pdf_text(content: bytes) -> str:
    doc = pdfium.PdfDocument(content)
    return "\n".join(page.get_textpage().get_text_range() for page in doc)


# --- formatting ---------------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("value", "currency", "locale", "expected"),
    [
        (Decimal("1234567.891"), "INR", "en-IN", "₹12,34,567.89"),
        (Decimal("-1234.5"), "USD", "en-US", "-$1,234.50"),
        (Decimal("1234.5"), "EUR", "de-DE", "1.234,50 €"),
        (Decimal("1234.5"), "AED", "en-AE", "AED 1,234.50"),
        (Decimal("1234.5"), "JPY", "en-US", "¥1,235"),
        (Decimal(0), "GBP", "en-GB", "£0.00"),
    ],
)
def test_format_money(value, currency, locale, expected):
    assert format_money(value, currency, locale) == expected


def test_format_money_without_symbols_uses_the_code():
    assert format_money(Decimal(1234), "INR", "en-IN", symbols=False) == "INR 1,234.00"


def test_excel_money_format():
    assert excel_money_format("INR", "en-IN") == '"₹"#,##0.00'
    assert excel_money_format("EUR", "fr-FR") == '#,##0.00 "€"'
    assert excel_money_format("SAR", "en-US") == '"SAR "#,##0.00'


def test_format_length_matches_the_app():
    assert format_length(to_mm("12.5", FT), LengthUnit.FT_IN) == "12'6\""
    assert format_length(to_mm("11.999", FT), LengthUnit.FT_IN) == "12'0\""
    assert format_length(Decimal(3000), LengthUnit.M) == "3 m"


# --- data ---------------------------------------------------------------------------------------


def test_not_ready_lists_the_blockers():
    doc = project(rooms=[room("Bedroom")])  # no size
    with pytest.raises(NotReady) as e:
        export_data(doc, "sqft")
    assert any("size" in b for b in e.value.blockers)


def test_shopping_list_has_one_item_per_tile_type():
    data = export_data(PRICED, "sqft")
    assert [s.ref for s in data.shopping] == ["T1", "T2"]
    assert {s.tile.tile.id for s in data.shopping} == {"t6", "t3"}
    assert sum(s.tile.boxes for s in data.shopping) == data.estimate.totals.boxes


# --- PDF ----------------------------------------------------------------------------------------


def test_report_pdf_has_table_totals_notice_costs_and_assumptions():
    data = export_data(PRICED, "sqft")
    text = pdf_text(render_report(data))
    assert NOTICE in text
    for r in data.estimate.rooms:
        assert r.name in text
    assert "Project total" in text
    assert str(data.estimate.totals.boxes) in text
    assert "Shopping list" in text
    for a in data.estimate.assumptions:
        assert a.title in text
    assert "Cost estimate" in text and "GST (18%)" in text and "Over budget" in text
    assert data.money(data.cost.total, symbols=False).split(" ")[1] in text


def test_report_without_prices_has_no_cost_section():
    data = export_data(UNPRICED, "sqm")
    assert not data.has_costs
    text = pdf_text(render_report(data))
    assert "Cost estimate" not in text and "Price / box" not in text
    assert "m²" in text


def test_shopping_list_pdf():
    data = export_data(PRICED, "sqft")
    text = pdf_text(render_shopping_list(data))
    assert "Ivory matt" in text and "Grey anti-skid" in text and NOTICE in text


# --- Excel --------------------------------------------------------------------------------------


def _roundup(value: float) -> int:
    return math.ceil(round(value, 9))


def test_workbook_uses_live_formulas_that_reproduce_the_app():
    data = export_data(PRICED, "sqft")
    wb = load_workbook(io.BytesIO(render_workbook(data)))
    assert wb.sheetnames == ["Rooms", "Shopping list", "Costs", "Assumptions"]
    ws = wb["Rooms"]
    assert NOTICE in ws["A2"].value

    est = {r.name: r for r in data.estimate.rooms}
    costs = {c.room_id: c for c in data.cost.rooms}
    n_rooms = len(data.estimate.rooms)
    for row in range(FIRST_ROW, FIRST_ROW + n_rooms):
        r = est[ws[f"B{row}"].value]
        assert ws[f"I{row}"].value == f"=G{row}+H{row}"
        assert ws[f"K{row}"].value == f"=ROUNDUP(I{row}/J{row},0)"
        assert ws[f"L{row}"].value == f"=K{row}*J{row}"
        # Evaluate the formulas by hand from the input cells: same boxes and cost as the app.
        required = ws[f"G{row}"].value + ws[f"H{row}"].value
        boxes = _roundup(required / ws[f"J{row}"].value)
        assert boxes == r.boxes
        assert round(boxes * ws[f"M{row}"].value, 2) == float(costs[r.room_id].cost)

    total = FIRST_ROW + n_rooms
    assert ws[f"B{total}"].value == "Project total"
    assert ws[f"K{total}"].value == f"=SUM(K{FIRST_ROW}:K{total - 1})"
    assert ws[f"N{total}"].value == f"=SUM(N{FIRST_ROW}:N{total - 1})"

    shop = wb["Shopping list"]
    assert shop[f"E{FIRST_ROW}"].value.startswith("=SUMIF(Rooms!")

    c = wb["Costs"]
    labels = [c.cell(r, 1).value for r in range(FIRST_ROW, c.max_row + 1)]
    for label in ("Tiles", "Delivery", "Fixing", "Handling", "GST", "Project total", "Status"):
        assert label in labels
    tax_row = FIRST_ROW + labels.index("GST")
    assert c[f"D{tax_row}"].value.startswith("=SUMIF(F")  # only the taxable items
    status = FIRST_ROW + labels.index("Status")
    assert c[f"E{status}"].value.startswith("=IF(")

    assumptions = wb["Assumptions"]
    titles = [assumptions.cell(r, 1).value for r in range(FIRST_ROW, assumptions.max_row + 1)]
    assert [a.title for a in data.estimate.assumptions] == titles[: len(data.estimate.assumptions)]


def test_workbook_without_costs_has_no_cost_sheet_or_columns():
    wb = load_workbook(io.BytesIO(render_workbook(export_data(UNPRICED, "sqft"))))
    assert "Costs" not in wb.sheetnames
    assert wb["Rooms"].cell(4, 13).value is None


# --- shopping list email ------------------------------------------------------------------------


def test_shopping_email_text_and_html_escape_user_text():
    doc = project()
    doc["name"] = "<b>Villa</b>"
    data = export_data(doc, "sqft")
    text = shopping_text(data, "Please quote")
    assert "Please quote" in text and NOTICE in text and "boxes" in text
    html = shopping_html(data, "A & B Tiling", "<script>x</script>")
    assert "<script>" not in html and "&lt;script&gt;" in html
    assert "&lt;b&gt;Villa&lt;/b&gt;" in html and "A &amp; B Tiling" in html


class _Settings:
    resend_api_key = "re_test"
    email_from = "Floor Planner <lists@example.com>"


@pytest.fixture
def resend(monkeypatch):
    """Route Resend calls to a fake; returns the list of requests it received."""
    calls: list[httpx.Request] = []
    reply = {"status": 200, "json": {"id": "msg_1"}}

    def handler(request: httpx.Request) -> httpx.Response:
        calls.append(request)
        return httpx.Response(reply["status"], json=reply["json"])

    real = httpx.AsyncClient
    monkeypatch.setattr(email, "get_settings", lambda: _Settings())
    monkeypatch.setattr(
        email.httpx, "AsyncClient", lambda **kw: real(transport=httpx.MockTransport(handler), **kw)
    )
    return calls, reply


async def test_send_email_posts_to_resend(resend):
    calls, _ = resend
    message_id = await email.send_email(
        to=["supplier@example.com"], subject="List", html="<p>x</p>", text="x", reply_to="me@x.com"
    )
    assert message_id == "msg_1"
    request = calls[0]
    assert str(request.url) == email.RESEND_URL
    assert request.headers["Authorization"] == "Bearer re_test"
    body = request.read().decode()
    assert '"from":"Floor Planner <lists@example.com>"' in body.replace(": ", ":")
    assert "supplier@example.com" in body and "me@x.com" in body


async def test_send_email_reports_resend_errors(resend):
    _, reply = resend
    reply.update(status=403, json={"message": "The example.com domain is not verified"})
    with pytest.raises(email.EmailError, match="not verified"):
        await email.send_email(to=["a@b.com"], subject="s", html="h", text="t")


async def test_send_email_needs_configuration(monkeypatch):
    class Off:
        resend_api_key = None
        email_from = None

    monkeypatch.setattr(email, "get_settings", lambda: Off())
    assert not email.email_enabled()
    with pytest.raises(email.EmailNotConfigured):
        await email.send_email(to=["a@b.com"], subject="s", html="h", text="t")
